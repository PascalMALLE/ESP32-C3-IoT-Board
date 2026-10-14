#!/usr/bin/env python3
"""
Convert a KiCad BOM CSV into NextPCB-compatible CSV and XLSX files.

Usage:
    python kicad_to_nextpcb_bom.py "my_kicad_bom.csv"

The input CSV should contain a reference/designator column. Optional columns:
MPN, Manufacturer, Footprint, Description, Procurement Type, Customer Note.
KiCad custom fields with these names are recommended.
"""
import csv
import re
import sys
from pathlib import Path
from collections import OrderedDict

HEADERS = [
    "Designator*", "Quantity*", "Manufacturer Part Number*",
    "Manufacturer", "Package/Footprint", "Description",
    "Procurement Type", "Customer Note"
]

ALIASES = {
    "ref": {"reference", "references", "ref", "refs", "designator", "designators", "reference(s)"},
    "value": {"value", "component value"},
    "mpn": {"mpn", "manufacturer part number", "manufacturerpartnumber", "part number", "partnumber"},
    "manufacturer": {"manufacturer", "mfr", "vendor"},
    "footprint": {"footprint", "package/footprint", "package", "package footprint"},
    "description": {"description", "comment", "comments"},
    "procurement": {"procurement type", "procurement", "sourcing", "assembly", "dnp"},
    "note": {"customer note", "customer notes", "note", "notes"},
}

def norm(s):
    return re.sub(r"[\s_\-]+", " ", (s or "").strip().lower())

def find_col(headers, field):
    normalized = {norm(h): h for h in headers}
    for alias in ALIASES[field]:
        if norm(alias) in normalized:
            return normalized[norm(alias)]
    # Handle headers that differ only by punctuation.
    compact = {re.sub(r"[^a-z0-9]", "", norm(h)): h for h in headers}
    for alias in ALIASES[field]:
        key = re.sub(r"[^a-z0-9]", "", norm(alias))
        if key in compact:
            return compact[key]
    return None

def split_refs(raw):
    # KiCad reference cells are usually comma-separated; accept semicolons too.
    refs = [x.strip() for x in re.split(r"[,;]", raw or "") if x.strip()]
    result = []
    for item in refs:
        # Also tolerate whitespace-separated references, but not ranges.
        if " " in item and not re.search(r"\s*-\s*", item):
            result.extend(x for x in item.split() if x)
        else:
            result.append(item)
    return result

def ref_key(ref):
    m = re.match(r"^(.*?)(\d+)([A-Za-z]*)$", ref)
    if m:
        return (m.group(1).lower(), int(m.group(2)), m.group(3).lower(), ref.lower())
    return (ref.lower(), -1, "", ref.lower())

def compress_refs(refs):
    """Compress consecutive same-prefix numeric refs, e.g. C1,C2,C3,C5 -> C1-C3,C5."""
    refs = sorted(set(refs), key=ref_key)
    parsed, other = {}, []
    for ref in refs:
        m = re.match(r"^(.*?)(\d+)$", ref)
        if not m:
            other.append(ref)
            continue
        prefix, number = m.group(1), int(m.group(2))
        parsed.setdefault(prefix, []).append((number, ref))
    out = []
    for prefix in sorted(parsed, key=str.lower):
        items = sorted(parsed[prefix])
        start = prev = items[0][0]
        start_ref = prev_ref = items[0][1]
        for number, ref in items[1:]:
            if number == prev + 1:
                prev, prev_ref = number, ref
            else:
                out.append(f"{start_ref}-{prev_ref}" if prev - start >= 2 else (f"{start_ref},{prev_ref}" if prev > start else start_ref))
                start = prev = number
                start_ref = prev_ref = ref
        out.append(f"{start_ref}-{prev_ref}" if prev - start >= 2 else (f"{start_ref},{prev_ref}" if prev > start else start_ref))
    out.extend(sorted(other, key=ref_key))
    # Flatten comma groups generated inside each prefix.
    return ",".join(out)

def read_csv(path):
    raw = path.read_text(encoding="utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(raw[:8192], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.DictReader(raw.splitlines(), dialect=dialect))
    if not rows:
        raise ValueError("Le CSV d'entrée est vide ou ne contient pas de ligne de composants.")
    return rows, list(rows[0].keys())

def main():
    if len(sys.argv) < 2:
        print('Usage: python kicad_to_nextpcb_bom.py "my_kicad_bom.csv"')
        raise SystemExit(2)
    source = Path(sys.argv[1]).expanduser()
    if not source.exists():
        raise FileNotFoundError(f"Fichier introuvable : {source}")
    rows, headers = read_csv(source)
    columns = {field: find_col(headers, field) for field in ALIASES}
    if not columns["ref"]:
        print("Colonnes détectées :", headers)
        raise ValueError("Colonne de références introuvable. Il faut une colonne Reference/References/Designator.")
    print("Correspondance des colonnes :")
    for field, col in columns.items():
        print(f"  {field:12} <- {col or '(absent; laissé vide)'}")

    # Group only truly equivalent items. Value is included internally to prevent
    # accidental grouping of different parts when no MPN is supplied.
    groups = OrderedDict()
    for row in rows:
        refs = split_refs(row.get(columns["ref"], "") if columns["ref"] else "")
        if not refs:
            continue
        get = lambda f: (row.get(columns[f], "") or "").strip() if columns.get(f) else ""
        mpn = get("mpn")
        manufacturer = get("manufacturer")
        footprint = get("footprint")
        description = get("description")
        procurement = get("procurement")
        note = get("note")
        value = get("value")
        key = (mpn, manufacturer, footprint, description, procurement, note, value)
        if key not in groups:
            groups[key] = {
                "refs": [], "mpn": mpn, "manufacturer": manufacturer,
                "footprint": footprint, "description": description,
                "procurement": procurement, "note": note
            }
        groups[key]["refs"].extend(refs)

    output_rows = []
    missing_mpn = []
    for group in groups.values():
        refs = sorted(set(group["refs"]), key=ref_key)
        if not group["mpn"]:
            missing_mpn.extend(refs)
        output_rows.append([
            compress_refs(refs), len(refs), group["mpn"], group["manufacturer"],
            group["footprint"], group["description"], group["procurement"], group["note"]
        ])
    output_rows.sort(key=lambda r: ref_key(r[0].split(",")[0].split("-")[0]))

    csv_path = source.with_name(source.stem + "_NextPCB.csv")
    xlsx_path = source.with_name(source.stem + "_NextPCB.xlsx")

    with csv_path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, quoting=csv.QUOTE_MINIMAL)
        writer.writerow(HEADERS)
        writer.writerows(output_rows)

    # XLSX with a clean import-ready BOM: header row followed immediately by parts.
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "BOM"
    ws.append(HEADERS)
    for row in output_rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = [24, 13, 31, 22, 31, 40, 20, 30]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    header_fill = PatternFill("solid", fgColor="17365D")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(xlsx_path)

    print(f"\\nOK : {len(output_rows)} lignes de composants générées.")
    print(f"CSV  : {csv_path}")
    print(f"XLSX : {xlsx_path}")
    if missing_mpn:
        print("\\nATTENTION : MPN manquant pour : " + ", ".join(sorted(set(missing_mpn), key=ref_key)))
        print("Renseigne le champ personnalisé 'MPN' dans KiCad avant l'envoi à NextPCB.")

if __name__ == "__main__":
    main()
